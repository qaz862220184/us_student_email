
import xlrd
import xlwt
import requests
#
# # 提取代理API接口，获取1个代理IP
# api_url = "http://dps.kdlapi.com/api/getdps/?orderid=979670906001794&num=1&pt=1&sep=1"
#
# # 获取API接口返回的代理IP
# proxy_ip = requests.get(api_url).text
# print(proxy_ip)
#
# # 用户名密码方式
# username = "sellerbdata"
# password = "lv6nv17v"
# proxies = {
#     "http": "http://%(user)s:%(pwd)s@%(proxy)s/" % {"user": username, "pwd": password, "proxy": proxy_ip},
#     "https": "http://%(user)s:%(pwd)s@%(proxy)s/" % {"user": username, "pwd": password, "proxy": proxy_ip}
# }
#
# #白名单方式（需提前设置白名单）
# # proxies = {
# #     "http": "http://%(proxy)s/" % {"proxy": proxy_ip},
# #     "https": "http://%(proxy)s/" % {"proxy": proxy_ip}
# # }
#
# # 要访问的目标网页
# target_url = "https://dev.kdlapi.com/testproxy"
#
# # 使用代理IP发送请求
# response = requests.get(target_url, proxies=proxies)
#
# # 获取页面内容
# if response.status_code == 200:
#     print(response.text)

'''
import openpyxl as xl
import os
headers = []
detail = []
user_detail = {'Full Name': 'Elena\xa0W\xa0Diehl', 'Gender': 'female', 'Title': 'Mrs.', 'Race': 'White', 'Birthday': '8/29/1970', 'Social Security Number': '768-38-0842', 'Street': '3435  Wilkinson Court', 'City': 'UTICA', 'State': 'IL', 'State Full': 'Illinois', 'Zip Code': '61373', 'Phone Number': '239-495-9205', 'Mobile Number': '847-314-1593', 'email': 'y2it72fo@temporary-mail.net','username': 'mnUOpaL7'}


def write_excel_file(folder_path,data):
    result_path = os.path.join(folder_path, "edu_zhuce.xlsx")
    print(result_path)
    print('***** 开始写入excel文件 ' + result_path + ' ***** \n')
    if os.path.exists(result_path):
        print('***** excel已存在，在表后添加数据 ' + result_path + ' ***** \n')
        workbook = xl.load_workbook(result_path)
        sheet = workbook.active
        for key in user_detail:
            detail.append(user_detail[key])
        sheet.append(detail)
        workbook.save(result_path)
    else:
        print('***** excel不存在，创建excel ' + result_path + ' ***** \n')
        workbook = xl.Workbook()
        workbook.save(result_path)

        sheet = workbook.active
        for item in user_detail:
            headers.append(item)
        sheet.append(headers)
        for key in user_detail:
            detail.append(user_detail[key])

        sheet.append(detail)
        workbook.save(result_path)
    print('***** 生成Excel文件 ' + result_path + ' ***** \n')


if __name__ == '__main__':
    write_excel_file("./")

'''
def screen_size():
    # 使用tkinter获取屏幕大小
    import tkinter
    tk = tkinter.Tk()
    width = tk.winfo_screenwidth()
    height = tk.winfo_screenheight()
    tk.quit()
    return width, height
import asyncio
import pyppeteer
from pyppeteer import launch
async def a():

    browser = await launch({
            #'userAgent': User_Agent1,
            'executablePath': pyppeteer.launcher.executablePath(),
            'headless': False,
            'dumpio': True,
            'autoClose': False,
            'args': [
                '--no-sandbox',
                "--start-maximized",
                '--disable-infobars',
                '--disable-extensions',
                '--hide-scrollbars',
                '--disable-bundled-ppapi-flash',
                '--mute-audio',
                '--no-sandbox',
                '--disable-setuid-sandbox',
                '--disable-gpu',
                '--enable-automation',
                #'--proxy-server={}'.format(proxy_ip)
                 #'--proxy-server=192.168.1.250:7890',
                '--proxy-server=192.168.1.239:24005'
            ]
        })

    page = await browser.newPage()
    width, height = screen_size()
    # print(width,height)
    await page.setViewport({  # 最大化窗口
        "width": width,
        "height": height
    })

    await page.evaluateOnNewDocument('() =>{ Object.defineProperties(navigator,'
                                         '{ webdriver:{ get: () => false } }) }')

    await page.goto('https://www.baidu.com/')


asyncio.get_event_loop().run_until_complete(a())

